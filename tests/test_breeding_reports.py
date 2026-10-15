from io import BytesIO
from datetime import date
from openpyxl import load_workbook
from sqlalchemy import select
from src.core.models import RoleModel,RolePermissionModel
from tests.conftest import login
from tests.test_animals import animal


def test_breeding_report_relations_filters_and_excel(api):
    client,factory=api;h=login(client)
    mother=animal(client,h,name='Luna',siniiga='MXM',birth_date='2020-01-01');male=animal(client,h,name='Trueno',siniiga='MXP',sex='male',birth_date='2019-01-01')
    r=client.post('/api/v1/breeding/matings',headers=h,json={'animal_id':mother['id'],'male_id':male['id'],'date':'2024-01-01'});assert r.status_code==201,r.text;mating=r.json()
    assert client.post('/api/v1/breeding/pregnancy-diagnoses',headers=h,json={'animal_id':mother['id'],'date':'2024-02-01','result':'pregnant','mating_event_id':mating['id']}).status_code==201
    r=client.post('/api/v1/breeding/births',headers=h,json={'animal_id':mother['id'],'father_id':male['id'],'date':'2024-06-01','birth_type_id':1,'offspring':[{'name':'=Estrella','siniiga':'MXC','sex':'female','birth_weight_kg':'3.25'},{'name':'Sin vida','sex':'male','alive':False}]});assert r.status_code==201,r.text;birth=r.json();child=next(c for c in birth['offspring'] if c['alive'])
    for f in ({'animal_id':male['id']},{'q':'Trueno'}):
        data=client.get('/api/v1/reports/breeding',headers=h,params=f).json();assert data['total']==1 and data['items'][0]['male']['id']==male['id']
    f={'kind':'pregnancy','result':'pregnant','date_from':'2024-02-01','date_to':'2024-02-01'}
    data=client.get('/api/v1/reports/breeding',headers=h,params=f).json();assert data['total']==1 and data['counts']=={'pregnant':1}
    assert client.get('/api/v1/reports/breeding',headers=h,params={**f,'result':'not_pregnant'}).json()['total']==0
    f={'kind':'birth','animal_id':child['animal']['id'],'limit':1}
    data=client.get('/api/v1/reports/breeding',headers=h,params=f).json();assert data['total']==1 and data['counts']=={'alive':1,'stillborn':1} and len(data['items'][0]['offspring'])==2
    assert data['items'][0]['offspring'][0]['birth_weight_kg'] in ('3.25',None)
    r=client.get('/api/v1/reports/breeding/export',headers=h,params=f);assert r.status_code==200
    book=load_workbook(BytesIO(r.content));assert book.sheetnames==['Partos','Crías']
    assert book['Partos']['E11'].value==1 and book['Partos']['F11'].value==1
    cells=[book['Crías'].cell(row,5) for row in (2,3)];c=next(c for c in cells if c.value=='=Estrella');assert c.data_type=='s'
    assert book['Crías'].cell(c.row,9).value==3.25 and book['Crías'].cell(c.row,9).data_type=='n'
    assert book['Crías'].cell(c.row,1).value.date()==date(2024,6,1)
    assert book['Partos'].freeze_panes=='A11' and book['Crías'].auto_filter.ref=='A1:I3'
    assert client.post('/api/v1/exits',headers=h,json={'animal_id':mother['id'],'exit_type':'sale','date':str(date.today()),'reason':'Venta'}).status_code==201
    assert client.get('/api/v1/reports/breeding',headers=h,params={'kind':'birth','q':'MXM'}).json()['total']==1
    options=client.get('/api/v1/reports/breeding/options',headers=h,params={'animal_q':'MXM'}).json();assert options['animals'][0]['id']==mother['id']


def test_breeding_report_permissions_and_validation(api):
    client,factory=api;h=login(client,'worker@example.com')
    for path in ('/reports/breeding','/reports/breeding/options','/reports/breeding/export'):
        assert client.get('/api/v1'+path,headers=h).status_code==403
        assert client.get('/api/v1'+path).status_code==401
    with factory.begin() as db:
        role=db.scalar(select(RoleModel).where(RoleModel.name=='worker'));db.add(RolePermissionModel(role_id=role.id,permission='breeding:read'))
    assert client.get('/api/v1/reports/breeding',headers=h).status_code==200
    assert client.get('/api/v1/reports/health',headers=h).status_code==403
    for f in ({'date_from':'2024-02-01','date_to':'2024-01-01'},{'kind':'birth','result':'pregnant'},{'animal_id':'bad'},{'kind':'bad'}):
        assert client.get('/api/v1/reports/breeding',headers=h,params=f).status_code==422
        assert client.get('/api/v1/reports/breeding/export',headers=h,params=f).status_code==422
