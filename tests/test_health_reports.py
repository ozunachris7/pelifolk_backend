from datetime import date,timedelta
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select
from src.core.models import RoleModel,RolePermissionModel
from tests.conftest import login
from tests.test_animals import animal
from tests.test_health import create_case,transition
from tests.test_treatments import product
from tests.test_programs import program


def test_health_case_report_filters_and_export(api):
    client,factory=api;h=login(client);a=animal(client,h,name='=Luna',siniiga='MX123',birth_date='2020-01-01')
    first=create_case(client,h,a['id'],reason='Cojera',notes='=Texto',start_date='2024-01-01').json()
    second=create_case(client,h,a['id'],reason='Revisión',start_date='2024-02-01').json()
    assert transition(client,h,second['id'],'open','closed',closed_date='2024-02-03').status_code==200
    f={'animal_id':a['id'],'date_from':'2024-02-01','date_to':'2024-02-01','status':'closed'}
    data=client.get('/api/v1/reports/health',headers=h,params=f).json()
    assert data['total']==1 and data['counts']=={'closed':1} and data['items'][0]['closed_date']=='2024-02-03'
    data=client.get('/api/v1/reports/health',headers=h,params={'q':'cojera','limit':1}).json()
    assert data['total']==1 and data['items'][0]['id']==first['id']
    r=client.get('/api/v1/reports/health/export',headers=h,params={'q':'cojera'});assert r.status_code==200
    sheet=load_workbook(BytesIO(r.content)).active
    assert sheet['A12'].value=='=Luna' and sheet['A12'].data_type=='s'
    assert sheet['G12'].value=='=Texto' and sheet['G12'].data_type=='s'
    assert sheet['E12'].value.date()==date(2024,1,1)
    assert sheet.freeze_panes=='A12' and sheet.auto_filter.ref=='A11:G12'
    assert client.get('/api/v1/reports/health/options',headers=h,params={'animal_q':'MX123'}).json()['animals'][0]['id']==a['id']


def test_health_schedule_report_and_dates(api):
    client,factory=api;h=login(client);today=date.today()
    a=animal(client,h,birth_date=str(today-timedelta(days=100)),siniiga='MX1')
    unknown=animal(client,h,name='Sin fecha',birth_date=None,siniiga='MX2')
    p=product(client,h).json();prog=program(client,h,p['id']).json()
    f={'kind':'schedule','status':'attention'}
    data=client.get('/api/v1/reports/health',headers=h,params=f).json()
    assert data['total']==2 and data['counts']=={'overdue':1,'missing_birth_date':1}
    data=client.get('/api/v1/reports/health',headers=h,params={**f,'date_from':str(today-timedelta(days=90)),'date_to':str(today)}).json()
    assert data['total']==1 and data['items'][0]['animal']['id']==a['id']
    sheet=load_workbook(BytesIO(client.get('/api/v1/reports/health/export',headers=h,params={**f,'animal_id':unknown['id']}).content)).active
    assert sheet['E12'].value=='Falta fecha de nacimiento' and sheet['G12'].value is None
    assert client.post('/api/v1/exits',headers=h,json={'animal_id':a['id'],'exit_type':'sale','date':str(today),'reason':'Venta'}).status_code==201
    assert client.get('/api/v1/reports/health',headers=h,params=f).json()['total']==1
    update={k:prog[k] for k in ('name','interval_days','minimum_age_days','notice_days','sex')};update.update(product_id=p['id'],active=False)
    assert client.put('/api/v1/health/programs/'+prog['id'],headers=h,json=update).status_code==200
    assert client.get('/api/v1/reports/health',headers=h,params=f).json()['total']==0


def test_health_report_access_and_validation(api):
    client,factory=api;worker=login(client,'worker@example.com')
    for path in ('/reports/health','/reports/health/options','/reports/health/export'):
        assert client.get('/api/v1'+path,headers=worker).status_code==403
        assert client.get('/api/v1'+path).status_code==401
    with factory.begin() as db:
        role=db.scalar(select(RoleModel).where(RoleModel.name=='worker'));db.add(RolePermissionModel(role_id=role.id,permission='health:read'))
    assert client.get('/api/v1/reports/health',headers=worker).status_code==200
    assert client.get('/api/v1/reports/animals',headers=worker).status_code==403
    for f in ({'date_from':'2024-02-01','date_to':'2024-01-01'},{'kind':'cases','status':'due'},{'kind':'schedule','status':'closed'},{'animal_id':'bad'}):
        assert client.get('/api/v1/reports/health',headers=worker,params=f).status_code==422
        assert client.get('/api/v1/reports/health/export',headers=worker,params=f).status_code==422
