from datetime import date
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select
from src.core.models import RoleModel,RolePermissionModel
from tests.conftest import login
from tests.test_finance import category


def test_finance_report_filtered_totals_period_balance_and_excel(api):
    client,factory=api;h=login(client)
    feed=category(client,h,'expense','feed');other=category(client,h,'expense','other');income=category(client,h,'income','animal_sales')
    supplier=client.post('/api/v1/finance/suppliers',headers=h,json={'name':'=Agropecuaria del Valle'}).json()
    for cat,amount,day,supplier_id in [(feed['id'],'100.10','2024-06-01',supplier['id']),(feed['id'],'50.05','2024-06-02',supplier['id']),(other['id'],'20.00','2024-06-02',None),(feed['id'],'900.00','2024-05-31',None)]:
        r=client.post('/api/v1/finance/expenses',headers=h,json={'category_id':cat,'amount':amount,'date':day,'supplier_id':supplier_id,'description':'=Compra registrada'});assert r.status_code==201,r.text
    assert client.post('/api/v1/finance/incomes',headers=h,json={'category_id':income['id'],'amount':'300.30','date':'2024-06-02','description':'Venta registrada'}).status_code==201
    assert client.put('/api/v1/finance/suppliers/'+supplier['id'],headers=h,json={'name':supplier['name'],'active':False}).status_code==200
    f={'date_from':'2024-06-01','date_to':'2024-06-02','category_id':feed['id'],'supplier_id':supplier['id'],'limit':1}
    data=client.get('/api/v1/reports/finance',headers=h,params=f).json()
    assert data['total']==2 and len(data['items'])==1 and data['filtered_total']=='150.15'
    assert data['summary']=={'income_total':'300.30','expense_total':'170.15','balance':'130.15','income_count':1,'expense_count':3}
    assert data['items'][0]['supplier']['active'] is False
    query=client.get('/api/v1/reports/finance',headers=h,params={**f,'q':'no match'}).json();assert query['filtered_total']=='0.00' and query['summary']==data['summary']
    r=client.get('/api/v1/reports/finance/export',headers=h,params=f);assert r.status_code==200
    book=load_workbook(BytesIO(r.content));assert book.sheetnames==['Resumen','Gastos']
    assert book['Resumen']['B8'].value==130.15 and book['Resumen']['B15'].value==150.15
    moves=book['Gastos'];assert moves.auto_filter.ref=='A1:F3' and moves.freeze_panes=='A2'
    assert moves['C2'].value==50.05 and moves['C2'].data_type=='n'
    assert moves['D2'].value=='=Agropecuaria del Valle' and moves['D2'].data_type=='s'
    assert moves['F2'].value=='=Compra registrada' and moves['F2'].data_type=='s'
    assert moves['A2'].value.date()==date(2024,6,2)
    options=client.get('/api/v1/reports/finance/options',headers=h,params={'supplier_q':'valle'}).json();assert options['suppliers']==[{'id':supplier['id'],'name':supplier['name'],'active':False}]
    inc=client.get('/api/v1/reports/finance',headers=h,params={'kind':'income','date_from':'2024-06-02','date_to':'2024-06-02'}).json();assert inc['filtered_total']=='300.30' and inc['total']==1


def test_finance_report_permissions_and_validation(api):
    client,factory=api;h=login(client,'worker@example.com')
    for path in ('/reports/finance','/reports/finance/options','/reports/finance/export'):
        assert client.get('/api/v1'+path,headers=h).status_code==403
        assert client.get('/api/v1'+path).status_code==401
    with factory.begin() as db:
        role=db.scalar(select(RoleModel).where(RoleModel.name=='worker'));db.add(RolePermissionModel(role_id=role.id,permission='finance:read'))
    assert client.get('/api/v1/reports/finance',headers=h).status_code==200
    assert client.get('/api/v1/reports/animals',headers=h).status_code==403
    for f in ({'date_from':'2024-02-01','date_to':'2024-01-01'},{'category_id':0},{'kind':'income','supplier_id':'00000000-0000-4000-8000-000000000001'},{'kind':'bad'}):
        assert client.get('/api/v1/reports/finance',headers=h,params=f).status_code==422
        assert client.get('/api/v1/reports/finance/export',headers=h,params=f).status_code==422
