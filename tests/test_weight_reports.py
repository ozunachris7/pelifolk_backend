from datetime import date
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select
from src.core.models import RoleModel, RolePermissionModel, EventModel
from tests.conftest import login
from tests.test_animals import animal


def test_weight_report_previous_before_range_pagination_and_excel(api):
    client, factory = api; h = login(client)
    a = animal(client, h, name='=Luna', siniiga='MX123', birth_date='2020-01-01')
    b = animal(client, h, name='Trueno', birth_date='2020-01-01')
    ids = []
    for animal_id, day, weight in [(a['id'],'2024-01-01','30'), (a['id'],'2024-01-11','32.50'), (a['id'],'2024-01-21','31.50'), (b['id'],'2024-01-11','40')]:
        r = client.post('/api/v1/weights', headers=h, json={'animal_id':animal_id,'date':day,'weight_kg':weight,'notes':'=Pesaje registrado'})
        assert r.status_code == 201, r.text
        ids.append(r.json()['id'])
    f = {'date_from':'2024-01-11','date_to':'2024-01-21','limit':1}
    r = client.get('/api/v1/reports/weights', headers=h, params=f); assert r.status_code == 200, r.text
    page = r.json()
    assert page['total'] == 3 and len(page['items']) == 1 and page['animal_count'] == 2
    assert page['comparable_count'] == 2 and page['loss_count'] == 1
    last = page['items'][0]
    assert last['previous_date'] == '2024-01-11' and last['days'] == 10
    assert last['gain_kg'] == '-1.00' and last['daily_gain_kg'] == '-0.100'
    filtered = client.get('/api/v1/reports/weights', headers=h, params={**f,'animal_id':a['id'],'date_to':'2024-01-11'}).json()
    assert filtered['total'] == 1
    assert filtered['items'][0]['previous_date'] == '2024-01-01'
    assert filtered['items'][0]['gain_kg'] == '2.50' and filtered['items'][0]['daily_gain_kg'] == '0.250'
    first = client.get('/api/v1/reports/weights', headers=h, params={'animal_id':b['id']}).json()['items'][0]
    assert first['previous_date'] is None and first['gain_kg'] is None and first['daily_gain_kg'] is None
    assert client.get('/api/v1/reports/weights', headers=h, params={'q':'%'}).json()['total'] == 0
    assert client.get('/api/v1/reports/weights', headers=h, params={'q':'mx123'}).json()['total'] == 3
    assert client.post('/api/v1/exits', headers=h, json={'animal_id':a['id'],'exit_type':'sale','date':'2024-01-22','reason':'Venta registrada'}).status_code == 201
    assert any(x['id'] == a['id'] for x in client.get('/api/v1/reports/weights/options',headers=h).json()['animals'])
    export = client.get('/api/v1/reports/weights/export', headers=h, params=f); assert export.status_code == 200
    book = load_workbook(BytesIO(export.content)); assert book.sheetnames == ['Resumen','Pesajes']
    sheet = book['Pesajes']; assert sheet.auto_filter.ref == 'A1:K4' and sheet.freeze_panes == 'A2'
    assert sheet['H2'].value == -1 and sheet['I2'].value == -0.1 and sheet['I2'].data_type == 'n'
    luna = next(r for r in sheet.iter_rows(min_row=2) if r[0].value.date()==date(2024,1,11) and r[1].value=='=Luna')
    assert luna[1].data_type == 's' and luna[10].data_type == 's' and luna[4].value.date() == date(2024,1,1)
    # Replaced events must not contribute to either the report or its previous measurement.
    with factory.begin() as db: db.get(EventModel,ids[1]).replaced_event_id = ids[0]
    filtered = client.get('/api/v1/reports/weights',headers=h,params={'animal_id':a['id'],'date_to':'2024-01-11'}).json()
    assert filtered['total'] == 1 and filtered['items'][0]['previous_date'] is None


def test_weight_report_permissions_and_validation(api):
    client,factory = api;h = login(client,'worker@example.com')
    for path in ('/reports/weights','/reports/weights/options','/reports/weights/export'):
        assert client.get('/api/v1'+path,headers=h).status_code == 403
        assert client.get('/api/v1'+path).status_code == 401
    with factory.begin() as db:
        role = db.scalar(select(RoleModel).where(RoleModel.name=='worker'))
        db.add(RolePermissionModel(role_id=role.id,permission='weights:read'))
    assert client.get('/api/v1/reports/weights',headers=h).status_code == 200
    assert client.get('/api/v1/reports/weights/export',headers=h).status_code == 200
    assert client.get('/api/v1/reports/animals',headers=h).status_code == 403
    for params in ({'date_from':'2024-02-01','date_to':'2024-01-01'},{'animal_id':'bad'},{'q':'x'*101}):
        for path in ('/reports/weights','/reports/weights/export'):
            assert client.get('/api/v1'+path,headers=h,params=params).status_code == 422
    assert client.get('/api/v1/reports/weights',headers=h,params={'limit':101}).status_code == 422
