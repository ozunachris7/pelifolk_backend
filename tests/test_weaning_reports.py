from datetime import date
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select
from src.core.models import RoleModel, RolePermissionModel
from tests.conftest import login
from tests.test_animals import animal
from tests.test_weanings import wean
from tests.test_weights import post as weigh


def test_report_filters_pagination_historical_data_and_excel(api):
    client,_=api; h=login(client)
    mother=animal(client,h,name='=Luna',birth_date='2020-01-01')
    child=animal(client,h,name='=Estrella',siniiga='MX123',birth_date='2024-01-01',mother_id=mother['id'])
    unknown=animal(client,h,name='Nube')
    assert weigh(client,h,child['id'],'2024-03-01','18.50',notes='Notas del pesaje').status_code==201
    assert wean(client,h,child,notes='=Notas del destete').status_code==201
    assert wean(client,h,unknown,date='2024-03-02',weight_kg='20').status_code==201
    other=animal(client,h,name='Otra madre',birth_date='2020-01-01')
    assert client.put('/api/v1/animals/'+child['id'],headers=h,json={'name':'=Estrella','sex':'female','siniiga':'MX123','birth_date':'2024-01-02','mother_id':other['id']}).status_code==200
    assert client.post('/api/v1/exits',headers=h,json={'animal_id':child['id'],'date':'2024-03-03','exit_type':'sale','reason':'Venta'}).status_code==201
    page=client.get('/api/v1/reports/weanings',headers=h,params={'limit':1}).json()
    assert page['total']==2 and len(page['items'])==1 and page['items'][0]['age_days'] is None
    second=client.get('/api/v1/reports/weanings',headers=h,params={'offset':1,'limit':1}).json()['items'][0]
    assert second['age_days']==60 and second['mother']['id']==mother['id']
    assert second['birth_date_snapshot']=='2024-01-01' and second['notes']=='=Notas del destete'
    f={'date_from':'2024-03-01','date_to':'2024-03-01','q':'mx123','animal_id':child['id']}
    filtered=client.get('/api/v1/reports/weanings',headers=h,params=f).json()
    assert filtered['total']==1 and filtered['items'][0]['animal']['id']==child['id']
    assert client.get('/api/v1/reports/weanings',headers=h,params={'q':'%'}).json()['total']==0
    options=client.get('/api/v1/reports/weanings/options',headers=h,params={'animal_q':'estrella'}).json()['animals']
    assert [a['id'] for a in options]==[child['id']]
    export=client.get('/api/v1/reports/weanings/export',headers=h,params={'limit':1})
    assert export.status_code==200 and export.headers['cache-control']=='no-store'
    book=load_workbook(BytesIO(export.content)); assert book.sheetnames==['Resumen','Destetes']
    sheet=book['Destetes']; assert sheet.max_row==3 and sheet.auto_filter.ref=='A1:J3' and sheet.freeze_panes=='A2'
    assert sheet['H2'].value is None
    assert sheet['A3'].value.date()==date(2024,3,1) and sheet['F3'].value.date()==date(2024,1,1)
    assert sheet['G3'].value==18.5 and sheet['G3'].data_type=='n' and sheet['H3'].value==60
    for address in ('B3','D3','J3'): assert sheet[address].value.startswith('=') and sheet[address].data_type=='s'
    book=load_workbook(BytesIO(client.get('/api/v1/reports/weanings/export',headers=h,params=f).content))
    assert book['Destetes'].max_row==2


def test_report_empty_permissions_and_invalid_filters(api):
    client,factory=api; h=login(client,'worker@example.com')
    endpoints=('/reports/weanings','/reports/weanings/options','/reports/weanings/export')
    for endpoint in endpoints:
        assert client.get('/api/v1'+endpoint,headers=h).status_code==403
        assert client.get('/api/v1'+endpoint).status_code==401
    with factory.begin() as db:
        role=db.scalar(select(RoleModel).where(RoleModel.name=='worker'))
        db.add(RolePermissionModel(role_id=role.id,permission='weights:read'))
    assert client.get('/api/v1/reports/weanings',headers=h).json()['total']==0
    export=client.get('/api/v1/reports/weanings/export',headers=h)
    assert export.status_code==200 and load_workbook(BytesIO(export.content))['Destetes'].max_row==1
    assert client.get('/api/v1/reports/weanings/options',headers=h).json()=={'animals':[]}
    assert client.get('/api/v1/reports/animals',headers=h).status_code==403
    for params in ({'date_from':'2024-04-01','date_to':'2024-03-01'},{'animal_id':'bad'},{'q':'x'*101}):
        for endpoint in ('/reports/weanings','/reports/weanings/export'):
            assert client.get('/api/v1'+endpoint,headers=h,params=params).status_code==422
    assert client.get('/api/v1/reports/weanings',headers=h,params={'limit':101}).status_code==422
