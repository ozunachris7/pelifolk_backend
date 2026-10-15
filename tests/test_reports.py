from io import BytesIO
from datetime import date
from openpyxl import load_workbook
from tests.conftest import login
from tests.test_animals import animal
from tests.test_organization import group, assign


def test_animal_report_filters_and_excel(api):
    client,factory=api;h=login(client)
    breed=client.post('/api/v1/breeds',headers=h,json={'name':'Pelibuey'}).json()
    a=animal(client,h,name='=HYPERLINK("bad")',birth_date='2020-01-01',breeds=[{'breed_id':breed['id'],'percentage':100}])
    male=animal(client,h,name='Trueno',sex='male',siniiga='MX02',birth_date='2020-01-01')
    g=group(client,h,'Reproductoras');assert assign(client,h,g['id'],[a['id']]).status_code==201
    params={'breed_id':breed['id'],'group_id':g['id'],'sex':'female'}
    data=client.get('/api/v1/reports/animals',headers=h,params=params).json()
    assert data['total']==1 and data['female']==1 and data['male']==0
    assert data['items'][0]['group']['name']=='Reproductoras'
    assert data['items'][0]['age_months']>=72
    assert client.get('/api/v1/reports/animals',headers=h,params={'sex':'male'}).json()['total']==1
    assert client.get('/api/v1/reports/animals',headers=h,params={'q':'%'}).json()['total']==0
    r=client.get('/api/v1/reports/animals/export',headers=h,params=params)
    assert r.status_code==200 and '.xlsx' in r.headers['content-disposition']
    sheet=load_workbook(BytesIO(r.content)).active
    assert sheet.freeze_panes=='A9' and sheet.auto_filter.ref=='A8:I9'
    assert sheet['B9'].value==a['name'] and sheet['B9'].data_type=='s'
    assert sheet['I9'].value=='Reproductoras' and sheet['E9'].value==data['items'][0]['age_months']
    assert client.post('/api/v1/exits',headers=h,json={'animal_id':male['id'],'exit_type':'sale','date':date.today().isoformat(),'reason':'Venta'}).status_code==201
    assert client.get('/api/v1/reports/animals',headers=h).json()['total']==1
    exits=client.get('/api/v1/reports/animals',headers=h,params={'status':'exited'}).json()
    assert exits['total']==1 and exits['items'][0]['status']=='sale'
    assert client.get('/api/v1/reports/animals',headers=h,params={'status':'all','limit':1}).json()['total']==2
    all_sheet=load_workbook(BytesIO(client.get('/api/v1/reports/animals/export',headers=h,params={'status':'all'}).content)).active
    assert all_sheet['B12'].value==2


def test_report_access_and_validation(api):
    client,factory=api;h=login(client,'worker@example.com')
    for path in ('/reports/animals','/reports/animals/options','/reports/animals/export'):
        assert client.get('/api/v1'+path,headers=h).status_code==403
        assert client.get('/api/v1'+path).status_code==401
    owner=login(client)
    for params in ({'breed_id':0},{'group_id':'bad'},{'status':'invalid'},{'limit':101}):
        assert client.get('/api/v1/reports/animals',headers=owner,params=params).status_code==422
