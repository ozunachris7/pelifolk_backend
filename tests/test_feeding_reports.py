from datetime import date
from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select
from src.core.models import UnitModel, FeedingModel, AnimalGroupModel, RoleModel, RolePermissionModel
from tests.conftest import login
from tests.test_animals import animal
from tests.test_feeding import feed_setup


def test_feeding_report_units_historical_prices_filters_and_excel(api):
    client,factory=api;h=login(client);feed=feed_setup(client,factory,h)
    a=animal(client,h,birth_date='2020-01-01',siniiga='MX123')
    for start,cost in [('2024-01-01','10.1250'),('2024-02-01','12.5000')]:
        assert client.post(f"/api/v1/feeding/feeds/{feed['id']}/prices",headers=h,json={'effective_from':start,'unit_cost':cost}).status_code==201
    with factory.begin() as db:db.add(UnitModel(id=3,code='bag'))
    bag=client.post('/api/v1/feeding/feeds',headers=h,json={'name':'=Concentrado','unit_id':3,'initial_price':{'effective_from':'2024-01-01','unit_cost':'100'}}).json()
    group=client.post('/api/v1/groups',headers=h,json={'name':'=Reproductoras'}).json()
    with factory.begin() as db:
        db.get(AnimalGroupModel,group['id']).active=False
        for day,quantity,feed_id,animal_id,group_id in [('2023-12-31','3',feed['id'],a['id'],None),('2024-01-31','1.25',feed['id'],a['id'],None),('2024-02-01','2',feed['id'],a['id'],None),('2024-02-01','1',bag['id'],None,group['id'])]:
            db.add(FeedingModel(date=date.fromisoformat(day),quantity=Decimal(quantity),feed_id=feed_id,animal_id=animal_id,group_id=group_id))
    data=client.get('/api/v1/reports/feeding',headers=h,params={'limit':1}).json()
    assert data['total']==4 and len(data['items'])==1 and data['total_cost']=='137.66' and data['unpriced_count']==1
    assert data['quantities']==[{'unit_code':'bag','quantity':'1.00'},{'unit_code':'kg','quantity':'6.25'}]
    filtered=client.get('/api/v1/reports/feeding',headers=h,params={'date_from':'2024-01-31','date_to':'2024-02-01','feed_id':feed['id'],'target_kind':'animal','target_id':a['id']}).json()
    assert filtered['total']==2 and filtered['total_cost']=='37.66'
    assert {r['price_from'] for r in filtered['items']}=={'2024-01-01','2024-02-01'}
    group_page=client.get('/api/v1/reports/feeding',headers=h,params={'target_kind':'group','target_id':group['id']}).json()
    assert group_page['total']==1 and group_page['total_cost']=='100.00'
    assert client.get('/api/v1/reports/feeding',headers=h,params={'q':'%'}).json()['total']==0
    options=client.get('/api/v1/reports/feeding/options',headers=h,params={'target_kind':'group','target_q':'reproductoras'}).json()
    assert options['targets']==[{'id':group['id'],'name':'=Reproductoras','siniiga':None}]
    response=client.get('/api/v1/reports/feeding/export',headers=h);assert response.status_code==200
    book=load_workbook(BytesIO(response.content));assert book.sheetnames==['Resumen','Consumos']
    assert book['Resumen']['B10'].value==137.66 and book['Resumen']['B11'].value==1
    assert book['Resumen']['B16'].value==1 and book['Resumen']['B17'].value==6.25
    sheet=book['Consumos'];assert sheet.auto_filter.ref=='A1:J5' and sheet.freeze_panes=='A2'
    rows=list(sheet.iter_rows(min_row=2,max_row=5))
    bag_row=next(r for r in rows if r[4].value=='=Concentrado')
    assert bag_row[2].value=='=Reproductoras' and bag_row[2].data_type=='s' and bag_row[4].data_type=='s'
    assert bag_row[6].data_type=='n' and bag_row[9].value==100
    legacy=next(r for r in rows if r[0].value.date()==date(2023,12,31))
    assert legacy[7].value is None and legacy[9].value is None
    january=next(r for r in rows if r[0].value.date()==date(2024,1,31))
    assert january[7].value==10.125 and january[9].value==12.66


def test_feeding_report_permissions_and_validation(api):
    client,factory=api;h=login(client,'worker@example.com')
    for path in ('/reports/feeding','/reports/feeding/options','/reports/feeding/export'):
        assert client.get('/api/v1'+path,headers=h).status_code==403
        assert client.get('/api/v1'+path).status_code==401
    with factory.begin() as db:
        role=db.scalar(select(RoleModel).where(RoleModel.name=='worker'));db.add(RolePermissionModel(role_id=role.id,permission='feeding:read'))
    assert client.get('/api/v1/reports/feeding',headers=h).status_code==200
    assert client.get('/api/v1/reports/animals',headers=h).status_code==403
    for params in ({'date_from':'2024-02-01','date_to':'2024-01-01'},{'target_id':'00000000-0000-4000-8000-000000000001'},{'target_kind':'bad'},{'feed_id':'bad'},{'q':'x'*101}):
        for path in ('/reports/feeding','/reports/feeding/export'):
            assert client.get('/api/v1'+path,headers=h,params=params).status_code==422
