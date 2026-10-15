from datetime import date, datetime, time
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from uuid import UUID

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from sqlalchemy import func, or_, select, type_coerce
from sqlalchemy.orm import Session

from src.core.database import get_db
from src.core.errors import ApplicationError
from src.core.models import AnimalModel, EventModel, EventWeightModel, UserModel
from src.features.auth.infrastructure.dependencies import require_permission
from src.features.weights.infrastructure.repository import SQLWeightRepository
from .routes import text_cell
from .breeding import style

router = APIRouter(prefix='/reports/weights', tags=['Reports'])


def filters(q: str = Query('', max_length=100), animal_id: UUID | None = None,
            date_from: date | None = None, date_to: date | None = None):
    if date_from and date_to and date_from > date_to:
        raise ApplicationError('Invalid report date range', 422)
    return dict(q=q.strip(), animal_id=str(animal_id) if animal_id else None,
                date_from=date_from, date_to=date_to)


def query_for(db, f):
    # Calculate the previous effective measurement before applying the report dates.
    order = (EventModel.date, EventModel.created_at, EventModel.id)
    query = SQLWeightRepository(db).query().with_only_columns(
        EventModel.id, EventModel.animal_id, AnimalModel.name, AnimalModel.siniiga,
        EventModel.date, EventWeightModel.weight_kg, EventModel.reason,
        UserModel.name.label('user_name'), UserModel.middle_name,
        UserModel.paternal_lastname, UserModel.maternal_lastname,
        type_coerce(func.lag(EventModel.date).over(partition_by=EventModel.animal_id, order_by=order), EventModel.date.type).label('previous_date'),
        type_coerce(func.lag(EventWeightModel.weight_kg).over(partition_by=EventModel.animal_id, order_by=order), EventWeightModel.weight_kg.type).label('previous_weight_kg'),
        maintain_column_froms=True)
    if f['animal_id']:
        query = query.where(EventModel.animal_id == f['animal_id'])
    if f['q']:
        query = query.where(or_(AnimalModel.name.icontains(f['q'], autoescape=True),
                                AnimalModel.siniiga.icontains(f['q'], autoescape=True)))
    history = query.subquery()
    result = select(history)
    if f['date_from']:
        result = result.where(history.c.date >= datetime.combine(f['date_from'], time.min))
    if f['date_to']:
        result = result.where(history.c.date <= datetime.combine(f['date_to'], time.max))
    return result.order_by(history.c.date.desc(), history.c.id)


def response(row):
    r = row._mapping
    day = r['date'].date()
    previous = r['previous_date'].date() if r['previous_date'] else None
    days = (day - previous).days if previous else None
    gain = r['weight_kg'] - r['previous_weight_kg'] if previous else None
    daily = (gain / Decimal(days)).quantize(Decimal('0.001'), rounding=ROUND_HALF_UP) if days and days > 0 else None
    return dict(id=r['id'], animal=dict(id=r['animal_id'], name=r['name'], siniiga=r['siniiga']),
                date=day.isoformat(), weight_kg=str(r['weight_kg']),
                previous_date=previous.isoformat() if previous else None,
                previous_weight_kg=str(r['previous_weight_kg']) if previous else None,
                days=days, gain_kg=str(gain) if gain is not None else None,
                daily_gain_kg=str(daily) if daily is not None else None,
                notes=r['reason'], recorded_by=' '.join(filter(None, [r['user_name'], r['middle_name'], r['paternal_lastname'], r['maternal_lastname']])))


def summary(rows):
    return dict(total=len(rows), animal_count=len({r['animal']['id'] for r in rows}),
                comparable_count=sum(r['daily_gain_kg'] is not None for r in rows),
                loss_count=sum(r['gain_kg'] is not None and Decimal(r['gain_kg']) < 0 for r in rows))


@router.get('/options')
def options(animal_q: str = Query('', max_length=100), actor=Depends(require_permission('weights:read')),
            db: Session = Depends(get_db, scope='function')):
    query = select(AnimalModel)
    if animal_q.strip():
        query = query.where(or_(AnimalModel.name.icontains(animal_q.strip(), autoescape=True),
                                AnimalModel.siniiga.icontains(animal_q.strip(), autoescape=True)))
    return dict(animals=[SQLWeightRepository.summary(a) for a in db.scalars(query.order_by(AnimalModel.name, AnimalModel.id).limit(100))])


@router.get('')
def report(f=Depends(filters), offset: int = Query(0, ge=0), limit: int = Query(25, ge=1, le=100),
           actor=Depends(require_permission('weights:read')), db: Session = Depends(get_db, scope='function')):
    rows = [response(r) for r in db.execute(query_for(db, f))]
    return dict(as_of=date.today(), date_from=f['date_from'], date_to=f['date_to'],
                items=rows[offset:offset + limit], **summary(rows), offset=offset, limit=limit)


@router.get('/export')
def export(f=Depends(filters), actor=Depends(require_permission('weights:read')), db: Session = Depends(get_db, scope='function')):
    today = date.today()
    rows = [response(r) for r in db.execute(query_for(db, f))]
    totals = summary(rows)
    book = Workbook(); sheet = book.active; sheet.title = 'Resumen'
    sheet.append(['Pelifolk · Pesajes y crecimiento']); sheet.append(['Fecha de consulta', today])
    for label, value in [('Desde', f['date_from']), ('Hasta', f['date_to'])]: sheet.append([label, value])
    for n in (2, 3, 4): sheet.cell(n, 2).number_format = 'dd/mm/yyyy'
    animal = db.get(AnimalModel, f['animal_id']) if f['animal_id'] else None
    for label, value in [('Búsqueda', f['q']), ('Animal', (animal.name or animal.siniiga or animal.id) if animal else 'Todos')]:
        sheet.append([label]); text_cell(sheet.cell(sheet.max_row, 2), value)
    for label, key in [('Pesajes', 'total'), ('Animales', 'animal_count'), ('Intervalos comparables', 'comparable_count'), ('Pesajes con pérdida', 'loss_count')]: sheet.append([label, totals[key]])
    sheet.append(['Cada pesaje se compara con el anterior efectivo del mismo animal, incluso fuera del rango.'])
    sheet.append(['Ganancia diaria = cambio de peso / días entre mediciones. El primer pesaje no tiene comparación.'])
    sheet.append(['Los cambios de peso no se suman entre animales. Valores negativos indican pérdida.'])
    sheet.column_dimensions['A'].width = 100; sheet.column_dimensions['B'].width = 35
    detail = book.create_sheet('Pesajes')
    detail.append(['Fecha', 'Animal', 'SINIIGA', 'Peso (kg)', 'Pesaje anterior', 'Peso anterior (kg)', 'Días', 'Cambio (kg)', 'Ganancia diaria (kg/día)', 'Registrado por', 'Notas'])
    for r in rows:
        detail.append([date.fromisoformat(r['date']), None, None, Decimal(r['weight_kg']),
                       date.fromisoformat(r['previous_date']) if r['previous_date'] else None,
                       Decimal(r['previous_weight_kg']) if r['previous_weight_kg'] is not None else None,
                       r['days'], Decimal(r['gain_kg']) if r['gain_kg'] is not None else None,
                       Decimal(r['daily_gain_kg']) if r['daily_gain_kg'] is not None else None, None, None])
        n = detail.max_row
        for col, value in [(2, r['animal']['name']), (3, r['animal']['siniiga']), (10, r['recorded_by']), (11, r['notes'])]: text_cell(detail.cell(n, col), value)
        for col in (1, 5): detail.cell(n, col).number_format = 'dd/mm/yyyy'
        for col in (4, 6, 8): detail.cell(n, col).number_format = '#,##0.00'
        detail.cell(n, 9).number_format = '#,##0.000'
    style(detail, 1)
    stream = BytesIO(); book.save(stream)
    return Response(stream.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename="pelifolk-pesajes-{today}.xlsx"', 'Cache-Control': 'no-store'})
