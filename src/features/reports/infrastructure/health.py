from datetime import date
from io import BytesIO
from typing import Literal
from uuid import UUID
from fastapi import APIRouter, Depends, Query, Response
from sqlalchemy import select, func, or_
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from src.core.database import get_db
from src.core.errors import ApplicationError
from src.core.models import AnimalModel, HealthCaseModel, HealthProgramModel, ProductModel
from src.features.auth.infrastructure.dependencies import require_permission
from src.features.health.infrastructure.program_repository import SQLProgramRepository
from .routes import text_cell

router=APIRouter(prefix='/reports/health',tags=['Reports'])
CASE_LABELS={'open':'Abierto','under_treatment':'En tratamiento','closed':'Cerrado'}
NOTICE_LABELS={'overdue':'Vencido','due':'Corresponde hoy','upcoming':'Próximo','scheduled':'Programado','missing_birth_date':'Falta fecha de nacimiento'}


def filters(kind: Literal['cases','schedule']='cases',q: str=Query('',max_length=100),
            status: Literal['all','open','under_treatment','closed','attention','overdue','due','upcoming','scheduled','missing_birth_date']='all',
            animal_id: UUID|None=None,program_id: UUID|None=None,date_from: date|None=None,date_to: date|None=None):
    if date_from and date_to and date_from>date_to:
        raise ApplicationError('Invalid report date range',422)
    allowed={'all',*CASE_LABELS} if kind=='cases' else {'all','attention',*NOTICE_LABELS}
    if status not in allowed:
        raise ApplicationError('Invalid report status',422)
    if kind=='cases' and program_id:
        raise ApplicationError('Program filter applies only to schedule',422)
    return dict(kind=kind,q=q.strip(),status=status,animal_id=str(animal_id) if animal_id else None,
                program_id=str(program_id) if program_id else None,date_from=date_from,date_to=date_to)


def case_query(f):
    query=select(HealthCaseModel,AnimalModel).join(AnimalModel,AnimalModel.id==HealthCaseModel.animal_id)
    if f['q']:
        query=query.where(or_(AnimalModel.name.icontains(f['q'],autoescape=True),AnimalModel.siniiga.icontains(f['q'],autoescape=True),HealthCaseModel.reason.icontains(f['q'],autoescape=True)))
    if f['animal_id']:query=query.where(HealthCaseModel.animal_id==f['animal_id'])
    if f['status']!='all':query=query.where(HealthCaseModel.status==f['status'])
    if f['date_from']:query=query.where(HealthCaseModel.start_date>=f['date_from'])
    if f['date_to']:query=query.where(HealthCaseModel.start_date<=f['date_to'])
    return query


def case_response(c,a):
    return dict(id=c.id,animal=dict(id=a.id,name=a.name,siniiga=a.siniiga),reason=c.reason,status=c.status,start_date=c.start_date,closed_date=c.closed_date,notes=c.notes)


def notices(db,f):
    repo=SQLProgramRepository(db)
    query=select(HealthProgramModel).join(ProductModel,ProductModel.id==HealthProgramModel.product_id).where(HealthProgramModel.active==True,ProductModel.active==True)
    if f['program_id']:query=query.where(HealthProgramModel.id==f['program_id'])
    rows=[]
    for p in db.scalars(query.order_by(HealthProgramModel.name,HealthProgramModel.id)):
        schedule=repo.schedule(p.id,f['q'],f['status'],0,2147483647)
        for row in schedule['items']:
            if f['animal_id'] and row['animal']['id']!=f['animal_id']:continue
            due=row['due_date']
            if f['date_from'] and (due is None or due<f['date_from']):continue
            if f['date_to'] and (due is None or due>f['date_to']):continue
            rows.append(dict(**row,id=p.id+':'+row['animal']['id'],program_id=p.id,program_name=p.name,product_name=schedule['program']['product']['name']))
    ranks={'overdue':0,'due':1,'upcoming':2,'missing_birth_date':3,'scheduled':4}
    rows.sort(key=lambda r:(ranks[r['status']],r['due_date'] or date.max,r['program_name'],r['animal']['id']))
    return rows


@router.get('/options')
def options(animal_q: str=Query('',max_length=100),actor=Depends(require_permission('health:read')),db:Session=Depends(get_db,scope='function')):
    animals=select(AnimalModel)
    if animal_q.strip():
        animals=animals.where(or_(AnimalModel.name.icontains(animal_q.strip(),autoescape=True),AnimalModel.siniiga.icontains(animal_q.strip(),autoescape=True)))
    programs=select(HealthProgramModel).join(ProductModel,ProductModel.id==HealthProgramModel.product_id).where(HealthProgramModel.active==True,ProductModel.active==True).order_by(HealthProgramModel.name)
    return dict(animals=[dict(id=a.id,name=a.name,siniiga=a.siniiga) for a in db.scalars(animals.order_by(AnimalModel.name,AnimalModel.id).limit(100))],programs=[dict(id=p.id,name=p.name) for p in db.scalars(programs)])


@router.get('')
def report(f=Depends(filters),offset:int=Query(0,ge=0),limit:int=Query(25,ge=1,le=100),actor=Depends(require_permission('health:read')),db:Session=Depends(get_db,scope='function')):
    if f['kind']=='cases':
        query=case_query(f)
        counts=dict(db.execute(query.with_only_columns(HealthCaseModel.status,func.count()).group_by(HealthCaseModel.status)).all())
        rows=[case_response(c,a) for c,a in db.execute(query.order_by(HealthCaseModel.start_date.desc(),HealthCaseModel.id).offset(offset).limit(limit))]
        total=sum(counts.values())
    else:
        all_rows=notices(db,f);total=len(all_rows);counts={}
        for row in all_rows:counts[row['status']]=counts.get(row['status'],0)+1
        rows=all_rows[offset:offset+limit]
    return dict(as_of=date.today(),kind=f['kind'],items=rows,total=total,counts=counts,offset=offset,limit=limit)


@router.get('/export')
def export(f=Depends(filters),actor=Depends(require_permission('health:read')),db:Session=Depends(get_db,scope='function')):
    today=date.today();book=Workbook();sheet=book.active;sheet.title='Casos' if f['kind']=='cases' else 'Programas sanitarios'
    sheet.append(['Pelifolk · Reporte de sanidad']);sheet.append(['Fecha de consulta',today]);sheet.cell(2,2).number_format='dd/mm/yyyy'
    labels=CASE_LABELS if f['kind']=='cases' else NOTICE_LABELS
    sheet.append(['Estado',labels.get(f['status'], 'Requieren atención' if f['status']=='attention' else 'Todos')])
    sheet.append(['Fechas aplicadas a','Inicio del caso' if f['kind']=='cases' else 'Próxima aplicación'])
    for label,value in [('Desde',f['date_from']),('Hasta',f['date_to'])]:sheet.append([label,value]);sheet.cell(sheet.max_row,2).number_format='dd/mm/yyyy'
    for label,value in [('Búsqueda',f['q']),('Animal',db.get(AnimalModel,f['animal_id']) if f['animal_id'] else None),('Programa',db.get(HealthProgramModel,f['program_id']) if f['program_id'] else None)]:
        if label=='Animal' and value:value=' · '.join(v for v in [value.name,value.siniiga] if v) or value.id
        elif label=='Programa' and value:value=value.name
        sheet.append([label]);text_cell(sheet.cell(sheet.max_row,2),value)
    sheet.append([])
    headers=['Animal','SINIIGA','Motivo','Estado','Inicio','Cierre','Observaciones'] if f['kind']=='cases' else ['Animal','SINIIGA','Programa','Producto','Estado','Última aplicación','Próxima aplicación','Días restantes']
    sheet.append(headers);header=sheet.max_row
    if f['kind']=='cases':
        rows=(case_response(c,a) for c,a in db.execute(case_query(f).order_by(HealthCaseModel.start_date.desc(),HealthCaseModel.id)))
    else:rows=notices(db,f)
    total=0
    for row in rows:
        if f['kind']=='cases':
            sheet.append([None,None,None,CASE_LABELS[row['status']],row['start_date'],row['closed_date'],None]);text_values={1:row['animal']['name'],2:row['animal']['siniiga'],3:row['reason'],7:row['notes']};dates=(5,6)
        else:
            sheet.append([None,None,None,None,NOTICE_LABELS[row['status']],row['last_application'],row['due_date'],row['days_remaining']]);text_values={1:row['animal']['name'],2:row['animal']['siniiga'],3:row['program_name'],4:row['product_name']};dates=(6,7)
        for col,value in text_values.items():text_cell(sheet.cell(sheet.max_row,col),value)
        for col in dates:sheet.cell(sheet.max_row,col).number_format='dd/mm/yyyy'
        total+=1
    last='G' if f['kind']=='cases' else 'H';sheet.auto_filter.ref=f'A{header}:{last}{sheet.max_row}';sheet.freeze_panes=f'A{header+1}'
    for cell in sheet[header]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='24583A')
    for col in 'ABCDEFGH':sheet.column_dimensions[col].width=26
    sheet.column_dimensions['C'].width=42;sheet.append([]);sheet.append(['Total de registros',total]);stream=BytesIO();book.save(stream)
    return Response(stream.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="pelifolk-sanidad-{f["kind"]}-{today}.xlsx"','Cache-Control':'no-store'})
