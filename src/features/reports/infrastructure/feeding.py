from datetime import date
from decimal import Decimal,ROUND_HALF_UP
from io import BytesIO
from typing import Literal
from uuid import UUID
from fastapi import APIRouter,Depends,Query,Response
from sqlalchemy import select,or_
from sqlalchemy.orm import Session
from openpyxl import Workbook
from src.core.database import get_db
from src.core.errors import ApplicationError
from src.core.models import FeedingModel,FeedModel,AnimalModel,AnimalGroupModel,UnitModel
from src.features.auth.infrastructure.dependencies import require_permission
from src.features.feeding.infrastructure.repository import SQLFeedingRepository
from src.features.feeding.domain.dto import FeedingResponse
from .routes import text_cell
from .breeding import style

router=APIRouter(prefix='/reports/feeding',tags=['Reports'])
UNIT_LABELS={'kg':'kg','lb':'lb','bag':'saco','bale':'paca','ml':'ml','tablet':'tableta','dose':'dosis'}


def filters(q:str=Query('',max_length=100),feed_id:UUID|None=None,target_kind:Literal['all','animal','group']='all',target_id:UUID|None=None,date_from:date|None=None,date_to:date|None=None):
    if date_from and date_to and date_from>date_to:raise ApplicationError('Invalid report date range',422)
    if target_id and target_kind=='all':raise ApplicationError('Choose a target kind before a target',422)
    return dict(q=q.strip(),feed_id=str(feed_id) if feed_id else None,target_kind=target_kind,target_id=str(target_id) if target_id else None,date_from=date_from,date_to=date_to)


def query_for(repo,f):
    query=repo.query()
    if f['q']:query=query.where(or_(FeedModel.name.icontains(f['q'],autoescape=True),AnimalModel.name.icontains(f['q'],autoescape=True),AnimalModel.siniiga.icontains(f['q'],autoescape=True),AnimalGroupModel.name.icontains(f['q'],autoescape=True)))
    if f['feed_id']:query=query.where(FeedingModel.feed_id==f['feed_id'])
    if f['target_kind']!='all':
        column=FeedingModel.animal_id if f['target_kind']=='animal' else FeedingModel.group_id
        query=query.where(column.is_not(None))
        if f['target_id']:query=query.where(column==f['target_id'])
    if f['date_from']:query=query.where(FeedingModel.date>=f['date_from'])
    if f['date_to']:query=query.where(FeedingModel.date<=f['date_to'])
    return query


def totals(db,repo,query):
    total=unpriced=0;cost=Decimal('0.00');units={}
    statement=query.join(UnitModel,UnitModel.id==FeedModel.unit_id).with_only_columns(FeedingModel.quantity,repo.price_expression(),UnitModel.code,maintain_column_froms=True)
    for quantity,price,code in db.execute(statement):
        total+=1;units[code]=units.get(code,Decimal('0.00'))+quantity
        if price is None:unpriced+=1
        else:cost+=(quantity*price).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)
    return dict(total=total,total_cost=str(cost),unpriced_count=unpriced,quantities=[dict(unit_code=code,quantity=str(units[code])) for code in sorted(units)])


@router.get('/options')
def options(feed_q:str=Query('',max_length=100),target_q:str=Query('',max_length=100),target_kind:Literal['animal','group']='animal',actor=Depends(require_permission('feeding:read')),db:Session=Depends(get_db,scope='function')):
    feeds=select(FeedModel).join(UnitModel,UnitModel.id==FeedModel.unit_id).add_columns(UnitModel.code)
    if feed_q.strip():feeds=feeds.where(FeedModel.name.icontains(feed_q.strip(),autoescape=True))
    model=AnimalModel if target_kind=='animal' else AnimalGroupModel;targets=select(model)
    if target_q.strip():
        conditions=[model.name.icontains(target_q.strip(),autoescape=True)]
        if target_kind=='animal':conditions.append(model.siniiga.icontains(target_q.strip(),autoescape=True))
        targets=targets.where(or_(*conditions))
    return dict(feeds=[dict(id=f.id,name=f.name,unit_code=code) for f,code in db.execute(feeds.order_by(FeedModel.name,FeedModel.id).limit(100))],targets=[dict(id=t.id,name=t.name,siniiga=t.siniiga if target_kind=='animal' else None) for t in db.scalars(targets.order_by(model.name,model.id).limit(100))])


@router.get('')
def report(f=Depends(filters),offset:int=Query(0,ge=0),limit:int=Query(25,ge=1,le=100),actor=Depends(require_permission('feeding:read')),db:Session=Depends(get_db,scope='function')):
    repo=SQLFeedingRepository(db);query=query_for(repo,f);summary=totals(db,repo,query)
    rows=[FeedingResponse.model_validate(repo.response(r)).model_dump(mode='json') for r in db.execute(query.order_by(FeedingModel.date.desc(),FeedingModel.id).offset(offset).limit(limit))]
    return dict(as_of=date.today(),date_from=f['date_from'],date_to=f['date_to'],items=rows,**summary,offset=offset,limit=limit)


@router.get('/export')
def export(f=Depends(filters),actor=Depends(require_permission('feeding:read')),db:Session=Depends(get_db,scope='function')):
    today=date.today();repo=SQLFeedingRepository(db);query=query_for(repo,f);summary=totals(db,repo,query);book=Workbook();sheet=book.active;sheet.title='Resumen'
    sheet.append(['Pelifolk · Reporte de alimentación']);sheet.append(['Fecha de consulta',today]);sheet.cell(2,2).number_format='dd/mm/yyyy'
    for label,value in [('Desde',f['date_from']),('Hasta',f['date_to'])]:sheet.append([label,value]);sheet.cell(sheet.max_row,2).number_format='dd/mm/yyyy'
    feed=db.get(FeedModel,f['feed_id']) if f['feed_id'] else None
    target=db.get(AnimalModel if f['target_kind']=='animal' else AnimalGroupModel,f['target_id']) if f['target_id'] else None
    for label,value in [('Búsqueda',f['q']),('Alimento',feed.name if feed else 'Todos'),('Destino',{'all':'Todos','animal':'Animales','group':'Lotes'}[f['target_kind']]),('Animal o lote',target.name if target else 'Todos')]:sheet.append([label]);text_cell(sheet.cell(sheet.max_row,2),value)
    sheet.append(['Consumos filtrados',summary['total']]);sheet.append(['Costo conocido (MXN)',Decimal(summary['total_cost'])]);sheet.cell(sheet.max_row,2).number_format='#,##0.00';sheet.append(['Consumos sin precio',summary['unpriced_count']]);sheet.append(['Los consumos sin precio no se incluyen en el costo conocido.'])
    sheet.append(['Las cantidades se suman por unidad, sin convertir kg, lb, sacos o pacas.']);sheet.append(['El consumo de un lote es del lote completo y no se distribuye por animal.']);sheet.append(['Unidad','Cantidad total'])
    for unit in summary['quantities']:sheet.append([UNIT_LABELS.get(unit['unit_code'],unit['unit_code']),Decimal(unit['quantity'])]);sheet.cell(sheet.max_row,2).number_format='#,##0.00'
    sheet.column_dimensions['A'].width=85;sheet.column_dimensions['B'].width=30
    detail=book.create_sheet('Consumos');detail.append(['Fecha','Tipo de destino','Animal o lote','SINIIGA','Alimento','Unidad','Cantidad','Precio unitario (MXN)','Precio vigente desde','Costo (MXN)'])
    for raw in db.execute(query.order_by(FeedingModel.date.desc(),FeedingModel.id)):
        row=repo.response(raw);target=row['animal'] or row['group'];detail.append([row['date'],'Animal' if row['animal'] else 'Lote',None,None,None,UNIT_LABELS.get(row['feed']['unit_code'],row['feed']['unit_code']),row['quantity'],row['unit_cost'],row['price_from'],row['cost']]);n=detail.max_row
        for col,value in [(3,target['name']),(4,target['siniiga']),(5,row['feed']['name'])]:text_cell(detail.cell(n,col),value)
        for col in (1,9):detail.cell(n,col).number_format='dd/mm/yyyy'
        for col in (7,10):detail.cell(n,col).number_format='#,##0.00'
        detail.cell(n,8).number_format='#,##0.0000'
    style(detail,1);detail.append([]);detail.append(['Costo conocido (MXN)',None,None,None,None,None,None,None,None,Decimal(summary['total_cost'])]);detail.cell(detail.max_row,10).number_format='#,##0.00';stream=BytesIO();book.save(stream)
    return Response(stream.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="pelifolk-alimentacion-{today}.xlsx"','Cache-Control':'no-store'})
