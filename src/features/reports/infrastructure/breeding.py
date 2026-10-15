from datetime import date, datetime, time
from io import BytesIO
from typing import Literal
from uuid import UUID
from fastapi import APIRouter, Depends, Query, Response
from sqlalchemy import select, exists, or_, func
from sqlalchemy.orm import Session, aliased
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from src.core.database import get_db
from src.core.errors import ApplicationError
from src.core.models import AnimalModel, EventModel, EventMatingModel, EventPregnancyModel, BirthOffspringModel
from src.features.auth.infrastructure.dependencies import require_permission
from src.features.breeding.infrastructure.repository import SQLBreedingRepository
from src.features.breeding.domain.dto import BreedingEvent
from .routes import text_cell

router=APIRouter(prefix='/reports/breeding',tags=['Reports'])
RESULTS={'pregnant':'Gestante','not_pregnant':'Vacía','uncertain':'Dudosa'}
BIRTHS={'natural':'Natural','assisted':'Asistido','cesarean':'Cesárea','other':'Otro'}
KINDS={'mating':'Montas','pregnancy':'Diagnósticos de gestación','birth':'Partos'}


def filters(kind:Literal['mating','pregnancy','birth']='mating',q:str=Query('',max_length=100),
            animal_id:UUID|None=None,result:Literal['pregnant','not_pregnant','uncertain']|None=None,
            date_from:date|None=None,date_to:date|None=None):
    if date_from and date_to and date_from>date_to:raise ApplicationError('Invalid report date range',422)
    if result and kind!='pregnancy':raise ApplicationError('Pregnancy result filter applies only to diagnoses',422)
    return dict(kind=kind,q=q.strip(),animal_id=str(animal_id) if animal_id else None,result=result,date_from=date_from,date_to=date_to)


def query_for(db,f):
    query=SQLBreedingRepository(db).query().where(EventModel.event_type_id=={'mating':5,'pregnancy':6,'birth':7}[f['kind']])
    if f['date_from']:query=query.where(EventModel.date>=datetime.combine(f['date_from'],time.min))
    if f['date_to']:query=query.where(EventModel.date<=datetime.combine(f['date_to'],time.max))
    if f['result']:query=query.where(exists(select(EventPregnancyModel.event_id).where(EventPregnancyModel.event_id==EventModel.id,EventPregnancyModel.result==f['result']).correlate(EventModel)))
    if f['animal_id']:
        animal_id=f['animal_id']
        query=query.where(or_(EventModel.animal_id==animal_id,
            exists(select(EventMatingModel.event_id).where(EventMatingModel.event_id==EventModel.id,EventMatingModel.male_id==animal_id)),
            exists(select(BirthOffspringModel.animal_id).where(BirthOffspringModel.event_id==EventModel.id,BirthOffspringModel.animal_id==animal_id))))
    if f['q']:
        a=aliased(AnimalModel)
        matching=select(a.id).where(or_(a.name.icontains(f['q'],autoescape=True),a.siniiga.icontains(f['q'],autoescape=True)))
        query=query.where(or_(EventModel.animal_id.in_(matching),
            exists(select(EventMatingModel.event_id).where(EventMatingModel.event_id==EventModel.id,EventMatingModel.male_id.in_(matching))),
            exists(select(BirthOffspringModel.animal_id).where(BirthOffspringModel.event_id==EventModel.id,BirthOffspringModel.animal_id.in_(matching)))))
    return query


def serialized(repo,event):
    return BreedingEvent.model_validate(repo.response(event)).model_dump(mode='json')


@router.get('/options')
def options(animal_q:str=Query('',max_length=100),actor=Depends(require_permission('breeding:read')),db:Session=Depends(get_db,scope='function')):
    query=select(AnimalModel)
    if animal_q.strip():query=query.where(or_(AnimalModel.name.icontains(animal_q.strip(),autoescape=True),AnimalModel.siniiga.icontains(animal_q.strip(),autoescape=True)))
    return dict(animals=[SQLBreedingRepository.summary(a) for a in db.scalars(query.order_by(AnimalModel.name,AnimalModel.id).limit(100))])


@router.get('')
def report(f=Depends(filters),offset:int=Query(0,ge=0),limit:int=Query(25,ge=1,le=100),actor=Depends(require_permission('breeding:read')),db:Session=Depends(get_db,scope='function')):
    query=query_for(db,f);total=db.scalar(select(func.count()).select_from(query.subquery()));counts={}
    if f['kind']=='pregnancy':
        counts=dict(db.execute(query.join(EventPregnancyModel,EventPregnancyModel.event_id==EventModel.id).with_only_columns(EventPregnancyModel.result,func.count()).group_by(EventPregnancyModel.result)).all())
    if f['kind']=='birth':
        counts=dict(db.execute(select(BirthOffspringModel.alive,func.count()).where(BirthOffspringModel.event_id.in_(query.with_only_columns(EventModel.id))).group_by(BirthOffspringModel.alive)).all())
        counts={'alive':counts.get(True,0),'stillborn':counts.get(False,0)}
    repo=SQLBreedingRepository(db)
    rows=[serialized(repo,e) for e in db.scalars(query.order_by(EventModel.date.desc(),EventModel.created_at.desc(),EventModel.id).offset(offset).limit(limit))]
    return dict(as_of=date.today(),kind=f['kind'],items=rows,total=total,counts=counts,offset=offset,limit=limit)


def style(sheet,header):
    from openpyxl.utils import get_column_letter
    sheet.freeze_panes=f'A{header+1}';sheet.auto_filter.ref=f'A{header}:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    for c in sheet[header]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='24583A')
    for n in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(n)].width=26


@router.get('/export')
def export(f=Depends(filters),actor=Depends(require_permission('breeding:read')),db:Session=Depends(get_db,scope='function')):
    today=date.today();book=Workbook();sheet=book.active;sheet.title={'mating':'Montas','pregnancy':'Gestación','birth':'Partos'}[f['kind']]
    sheet.append(['Pelifolk · Reporte de reproducción']);sheet.append(['Fecha de consulta',today]);sheet.cell(2,2).number_format='dd/mm/yyyy';sheet.append(['Registros',KINDS[f['kind']]])
    for label,value in [('Desde',f['date_from']),('Hasta',f['date_to'])]:sheet.append([label,value]);sheet.cell(sheet.max_row,2).number_format='dd/mm/yyyy'
    animal=db.get(AnimalModel,f['animal_id']) if f['animal_id'] else None
    for label,value in [('Búsqueda',f['q']),('Animal relacionado',' · '.join(v for v in [animal.name,animal.siniiga] if v) if animal else ''),('Resultado',RESULTS.get(f['result'],'Todos'))]:
        sheet.append([label]);text_cell(sheet.cell(sheet.max_row,2),value)
    sheet.append([])
    detail={'mating':['Macho','SINIIGA del macho'],'pregnancy':['Resultado','ID de monta vinculada'],'birth':['Tipo de parto','Crías vivas','Crías sin vida']}[f['kind']]
    headers=['Fecha','Hembra','SINIIGA de la hembra',*detail,'Observaciones','Registrado por'];sheet.append(headers);header=sheet.max_row
    children=None
    if f['kind']=='birth':
        children=book.create_sheet('Crías');children.append(['Fecha del parto','ID del parto','Madre','SINIIGA de la madre','Cría','SINIIGA de la cría','Sexo','Nació viva','Peso al nacer (kg)'])
    repo=SQLBreedingRepository(db);total=0;child_count=0
    for event in db.scalars(query_for(db,f).order_by(EventModel.date.desc(),EventModel.created_at.desc(),EventModel.id)):
        row=repo.response(event)
        if f['kind']=='mating':values=[row['male']['name'],row['male']['siniiga']]
        elif f['kind']=='pregnancy':values=[RESULTS[row['result']],row['mating_event_id']]
        else:
            values=[BIRTHS.get(row['birth_type'],row['birth_type']),sum(c['alive'] for c in row['offspring']),sum(not c['alive'] for c in row['offspring'])]
            for c in row['offspring']:
                children.append([row['date'],None,None,None,None,None,'Hembra' if c['animal']['sex']=='female' else 'Macho','Sí' if c['alive'] else 'No',c['birth_weight_kg']]);n=children.max_row
                for col,value in [(2,row['id']),(3,row['animal']['name']),(4,row['animal']['siniiga']),(5,c['animal']['name']),(6,c['animal']['siniiga'])]:text_cell(children.cell(n,col),value)
                children.cell(n,1).number_format='dd/mm/yyyy';children.cell(n,9).number_format='0.00';child_count+=1
        all_values=[row['date'],row['animal']['name'],row['animal']['siniiga'],*values,row['notes'],row['recorded_by']]
        sheet.append([None]*len(all_values));n=sheet.max_row
        for col,value in enumerate(all_values,1):
            if isinstance(value,(date,int)):sheet.cell(n,col).value=value
            else:text_cell(sheet.cell(n,col),value)
        sheet.cell(n,1).number_format='dd/mm/yyyy';total+=1
    style(sheet,header);sheet.append([]);sheet.append(['Total de registros',total])
    if children is not None:style(children,1);children.append([]);children.append(['Total de crías',child_count])
    stream=BytesIO();book.save(stream)
    return Response(stream.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="pelifolk-reproduccion-{f["kind"]}-{today}.xlsx"','Cache-Control':'no-store'})
