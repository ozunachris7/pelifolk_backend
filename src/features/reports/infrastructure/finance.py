from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Literal
from uuid import UUID
from fastapi import APIRouter, Depends, Query, Response
from sqlalchemy import select,func,or_
from sqlalchemy.orm import Session
from openpyxl import Workbook
from src.core.database import get_db
from src.core.errors import ApplicationError
from src.core.models import ExpenseModel,IncomeModel,SupplierModel,AnimalModel
from src.features.auth.infrastructure.dependencies import require_permission
from src.features.finance.infrastructure.repository import SQLFinanceRepository
from src.features.finance.domain.dto import MoneyResponse,FinanceSummary
from .routes import text_cell
from .breeding import style

router=APIRouter(prefix='/reports/finance',tags=['Reports'])
CATEGORY_LABELS={'animal_sales':'Venta de animales','feed':'Alimentación','veterinary':'Sanidad','transport':'Transporte','maintenance':'Mantenimiento','other':'Otros'}


def filters(kind:Literal['expense','income']='expense',q:str=Query('',max_length=100),category_id:int|None=Query(None,gt=0),
            supplier_id:UUID|None=None,date_from:date|None=None,date_to:date|None=None):
    if date_from and date_to and date_from>date_to:raise ApplicationError('Invalid report date range',422)
    if kind=='income' and supplier_id:raise ApplicationError('Supplier filter applies only to expenses',422)
    return dict(kind=kind,q=q.strip(),category_id=category_id,supplier_id=str(supplier_id) if supplier_id else None,date_from=date_from,date_to=date_to)


def query_for(repo,f):
    kind=f['kind'];model=ExpenseModel if kind=='expense' else IncomeModel;category=repo.category_model(kind)
    query=repo.range(repo.money_query(kind),model.date,f['date_from'],f['date_to'])
    if f['category_id']:query=query.where((model.expense_category_id if kind=='expense' else model.income_category_id)==f['category_id'])
    if f['supplier_id']:query=query.where(ExpenseModel.supplier_id==f['supplier_id'])
    if f['q']:
        conditions=[model.description.icontains(f['q'],autoescape=True),category.name.icontains(f['q'],autoescape=True)]
        conditions += [SupplierModel.name.icontains(f['q'],autoescape=True)] if kind=='expense' else [AnimalModel.name.icontains(f['q'],autoescape=True),AnimalModel.siniiga.icontains(f['q'],autoescape=True)]
        query=query.where(or_(*conditions))
    return query


def totals(db,query,model):
    amount,count=db.execute(query.with_only_columns(func.coalesce(func.sum(model.amount),0),func.count(model.id),maintain_column_froms=True)).one()
    return Decimal(amount).quantize(Decimal('0.01')),count


@router.get('/options')
def options(supplier_q:str=Query('',max_length=100),actor=Depends(require_permission('finance:read')),db:Session=Depends(get_db,scope='function')):
    repo=SQLFinanceRepository(db);query=select(SupplierModel)
    if supplier_q.strip():query=query.where(SupplierModel.name.icontains(supplier_q.strip(),autoescape=True))
    return dict(expense_categories=repo.categories('expense'),income_categories=repo.categories('income'),suppliers=[dict(id=s.id,name=s.name,active=s.active) for s in db.scalars(query.order_by(SupplierModel.name,SupplierModel.id).limit(100))])


@router.get('')
def report(f=Depends(filters),offset:int=Query(0,ge=0),limit:int=Query(25,ge=1,le=100),actor=Depends(require_permission('finance:read')),db:Session=Depends(get_db,scope='function')):
    repo=SQLFinanceRepository(db);model=ExpenseModel if f['kind']=='expense' else IncomeModel;query=query_for(repo,f);total,count=totals(db,query,model)
    items=[MoneyResponse.model_validate(repo.money_response(f['kind'],r)).model_dump(mode='json') for r in db.execute(query.order_by(model.date.desc(),model.id).offset(offset).limit(limit))]
    return dict(as_of=date.today(),date_from=f['date_from'],date_to=f['date_to'],kind=f['kind'],items=items,total=count,filtered_total=str(total),summary=FinanceSummary.model_validate(repo.summary(f['date_from'],f['date_to'])).model_dump(mode='json'),offset=offset,limit=limit)


@router.get('/export')
def export(f=Depends(filters),actor=Depends(require_permission('finance:read')),db:Session=Depends(get_db,scope='function')):
    today=date.today();repo=SQLFinanceRepository(db);model=ExpenseModel if f['kind']=='expense' else IncomeModel;query=query_for(repo,f)
    amount,count=totals(db,query,model);summary=repo.summary(f['date_from'],f['date_to']);book=Workbook();sheet=book.active;sheet.title='Resumen'
    sheet.append(['Pelifolk · Reporte de economía']);sheet.append(['Fecha de consulta',today]);sheet.cell(2,2).number_format='dd/mm/yyyy'
    for label,value in [('Desde',f['date_from']),('Hasta',f['date_to'])]:sheet.append([label,value]);sheet.cell(sheet.max_row,2).number_format='dd/mm/yyyy'
    sheet.append(['Balance: todos los movimientos del rango de fechas; sin filtros de lista.'])
    for label,key in [('Ingresos del período','income_total'),('Gastos del período','expense_total'),('Balance del período','balance')]:sheet.append([label,summary[key]]);sheet.cell(sheet.max_row,2).number_format='#,##0.00'
    sheet.append(['Moneda','MXN']);sheet.append(['Vista','Gastos' if f['kind']=='expense' else 'Ingresos'])
    category=db.get(repo.category_model(f['kind']),f['category_id']) if f['category_id'] else None;supplier=db.get(SupplierModel,f['supplier_id']) if f['supplier_id'] else None
    for label,value in [('Búsqueda',f['q']),('Categoría',CATEGORY_LABELS.get(category.name,category.name) if category else 'Todas'),('Proveedor',supplier.name if supplier else 'Todos')]:sheet.append([label]);text_cell(sheet.cell(sheet.max_row,2),value)
    sheet.append(['Movimientos filtrados',count]);sheet.append(['Importe filtrado',amount]);sheet.cell(sheet.max_row,2).number_format='#,##0.00';sheet.column_dimensions['A'].width=75;sheet.column_dimensions['B'].width=30
    moves=book.create_sheet('Gastos' if f['kind']=='expense' else 'Ingresos');moves.append(['Fecha','Categoría','Importe (MXN)','Proveedor' if f['kind']=='expense' else 'Animal vinculado','SINIIGA' if f['kind']=='income' else 'Estado del proveedor','Descripción'])
    if f['kind']=='income':
        for col,title in [(7,'Peso de venta (kg)'),(8,'Precio por kg (MXN)'),(9,'ID del pesaje')]:moves.cell(1,col,title)
    for raw in db.execute(query.order_by(model.date.desc(),model.id)):
        row=repo.money_response(f['kind'],raw);moves.append([row['date'],None,row['amount'],None,None,None]);n=moves.max_row
        if f['kind']=='expense':associated=row['supplier']['name'] if row['supplier'] else '';detail=('Activo' if row['supplier']['active'] else 'Inactivo') if row['supplier'] else ''
        else:associated=row['animal']['name'] if row.get('animal') else '';detail=row['animal']['siniiga'] if row.get('animal') else ''
        for col,value in [(2,CATEGORY_LABELS.get(row['category']['name'],row['category']['name'])),(4,associated),(5,detail),(6,row['description'])]:text_cell(moves.cell(n,col),value)
        if f['kind']=='income' and row.get('sale_pricing'):
            pricing=row['sale_pricing']
            moves.cell(n,7,pricing['weight_kg']).number_format='0.00'
            moves.cell(n,8,pricing['price_per_kg']).number_format='#,##0.00'
            text_cell(moves.cell(n,9),pricing['weight_event_id'])
        moves.cell(n,1).number_format='dd/mm/yyyy';moves.cell(n,3).number_format='#,##0.00'
    style(moves,1);moves.append([]);moves.append(['Total de movimientos',count]);moves.append(['Importe total (MXN)',None,amount]);moves.cell(moves.max_row,3).number_format='#,##0.00';stream=BytesIO();book.save(stream)
    return Response(stream.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="pelifolk-economia-{f["kind"]}-{today}.xlsx"','Cache-Control':'no-store'})
