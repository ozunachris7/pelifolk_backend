from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from src.core.database import get_db
from src.core.errors import ApplicationError
from src.core.models import AnimalModel, EventModel, EventWeaningModel
from src.features.auth.infrastructure.dependencies import require_permission
from src.features.weaning.repository import SQLWeaningRepository
from .breeding import style
from .routes import text_cell

router = APIRouter(prefix='/reports/weanings', tags=['Reports'])


def filters(q: str = Query('', max_length=100), animal_id: UUID | None = None,
            date_from: date | None = None, date_to: date | None = None):
    if date_from and date_to and date_from > date_to:
        raise ApplicationError('Invalid report date range', 422)
    return dict(q=q.strip(), animal_id=str(animal_id) if animal_id else None,
                date_from=date_from, date_to=date_to)


def query_for(db, f):
    query = SQLWeaningRepository(db).query()
    if f['q']:
        query = query.where(or_(AnimalModel.name.icontains(f['q'], autoescape=True),
                               AnimalModel.siniiga.icontains(f['q'], autoescape=True)))
    if f['animal_id']:
        query = query.where(EventWeaningModel.animal_id == f['animal_id'])
    if f['date_from']:
        query = query.where(EventModel.date >= datetime.combine(f['date_from'], time.min))
    if f['date_to']:
        query = query.where(EventModel.date <= datetime.combine(f['date_to'], time.max))
    return query.order_by(EventModel.date.desc(), EventModel.id)


@router.get('/options')
def options(animal_q: str = Query('', max_length=100), actor=Depends(require_permission('weights:read')),
            db: Session = Depends(get_db, scope='function')):
    query = select(AnimalModel).join(EventWeaningModel, EventWeaningModel.animal_id == AnimalModel.id)
    if animal_q.strip():
        query = query.where(or_(AnimalModel.name.icontains(animal_q.strip(), autoescape=True),
                               AnimalModel.siniiga.icontains(animal_q.strip(), autoescape=True)))
    return dict(animals=[SQLWeaningRepository.summary(a) for a in db.scalars(
        query.order_by(AnimalModel.name, AnimalModel.id).limit(100))])


@router.get('')
def report(f=Depends(filters), offset: int = Query(0, ge=0), limit: int = Query(25, ge=1, le=100),
           actor=Depends(require_permission('weights:read')), db: Session = Depends(get_db, scope='function')):
    query = query_for(db, f)
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery()))
    repo = SQLWeaningRepository(db)
    items = [repo.response(r) for r in db.execute(query.offset(offset).limit(limit))]
    return dict(as_of=date.today(), date_from=f['date_from'], date_to=f['date_to'],
                items=items, total=total, offset=offset, limit=limit)


@router.get('/export')
def export(f=Depends(filters), actor=Depends(require_permission('weights:read')),
           db: Session = Depends(get_db, scope='function')):
    today = date.today()
    book = Workbook()
    summary = book.active
    summary.title = 'Resumen'
    summary.append(['Pelifolk · Destetes'])
    for label, value in [('Fecha de consulta', today), ('Desde', f['date_from']), ('Hasta', f['date_to'])]:
        summary.append([label, value])
        summary.cell(summary.max_row, 2).number_format = 'dd/mm/yyyy'
    animal = db.get(AnimalModel, f['animal_id']) if f['animal_id'] else None
    for label, value in [('Búsqueda', f['q']), ('Animal', (animal.name or animal.siniiga or animal.id) if animal else 'Todos')]:
        summary.append([label])
        text_cell(summary.cell(summary.max_row, 2), value)
    summary.append(['La madre y el nacimiento corresponden a los datos guardados al registrar el destete.'])
    summary.append(['Una edad vacía indica que no había fecha de nacimiento registrada; no equivale a cero.'])
    summary.append(['Incluye animales que posteriormente salieron del hato.'])
    summary.column_dimensions['A'].width = 110
    summary.column_dimensions['B'].width = 35
    detail = book.create_sheet('Destetes')
    detail.append(['Fecha de destete', 'Animal', 'SINIIGA', 'Madre al registrar', 'SINIIGA de la madre',
                   'Nacimiento al registrar', 'Peso al destete (kg)', 'Edad al destete (días)', 'Registrado por', 'Observaciones'])
    repo = SQLWeaningRepository(db)
    total = 0
    for row in db.execute(query_for(db, f)):
        r = repo.response(row)
        mother = r['mother'] or {}
        detail.append([r['date'], None, None, None, None, r['birth_date_snapshot'],
                       Decimal(r['weight_kg']), r['age_days'], None, None])
        n = detail.max_row
        for col, value in [(2, r['animal']['name']), (3, r['animal']['siniiga']),
                           (4, mother.get('name')), (5, mother.get('siniiga')),
                           (9, r['recorded_by']), (10, r['notes'])]:
            text_cell(detail.cell(n, col), value)
        for col in (1, 6):
            detail.cell(n, col).number_format = 'dd/mm/yyyy'
        detail.cell(n, 7).number_format = '#,##0.00'
        detail.cell(n, 8).number_format = '0'
        total += 1
    summary.append(['Destetes filtrados', total])
    style(detail, 1)
    stream = BytesIO()
    book.save(stream)
    return Response(stream.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename="pelifolk-destetes-{today}.xlsx"', 'Cache-Control': 'no-store'})
