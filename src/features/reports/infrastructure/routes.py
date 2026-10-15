from datetime import date
from io import BytesIO
from typing import Literal
from uuid import UUID
from fastapi import APIRouter, Depends, Query, Response
from sqlalchemy import select, exists, or_, func
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from src.core.database import get_db
from src.core.herd import has_exited, exit_state
from src.core.models import AnimalModel, AnimalBreedModel, BreedModel, AnimalGroupModel, GroupAnimalModel
from src.features.auth.infrastructure.dependencies import require_permission

router=APIRouter(prefix='/reports',tags=['Reports'])


def filters(q: str=Query('',max_length=100), status: Literal['active','exited','all']='active',
            sex: Literal['male','female']|None=None, breed_id: int|None=Query(None,gt=0),group_id: UUID|None=None):
    return dict(q=q.strip(),status=status,sex=sex,breed_id=breed_id,group_id=str(group_id) if group_id else None)


def membership(day):
    return (GroupAnimalModel.start_date<=day,or_(GroupAnimalModel.end_date.is_(None),GroupAnimalModel.end_date>=day))


def query_for(f,day):
    query=select(AnimalModel)
    if f['q']:
        query=query.where(or_(AnimalModel.name.icontains(f['q'],autoescape=True),AnimalModel.siniiga.icontains(f['q'],autoescape=True)))
    if f['status']!='all':
        exited=has_exited(AnimalModel.id,day)
        query=query.where(~exited if f['status']=='active' else exited)
    if f['sex']:
        query=query.where(AnimalModel.sex==f['sex'])
    if f['breed_id']:
        query=query.where(exists(select(AnimalBreedModel.animal_id).where(AnimalBreedModel.animal_id==AnimalModel.id,AnimalBreedModel.breed_id==f['breed_id'])))
    if f['group_id']:
        query=query.where(~has_exited(AnimalModel.id,day),exists(select(GroupAnimalModel.animal_id).where(GroupAnimalModel.animal_id==AnimalModel.id,GroupAnimalModel.group_id==f['group_id'],*membership(day))))
    return query


def response(db,animal,day):
    status,exit_date=exit_state(db,animal.id)
    group=None
    if status=='active':
        row=db.execute(select(AnimalGroupModel).join(GroupAnimalModel,GroupAnimalModel.group_id==AnimalGroupModel.id).where(GroupAnimalModel.animal_id==animal.id,*membership(day)).order_by(GroupAnimalModel.start_date.desc(),AnimalGroupModel.id).limit(1)).scalar()
        if row: group=dict(id=row.id,name=row.name)
    born=animal.birth_date
    months=None if born is None else (day.year-born.year)*12+day.month-born.month-(day.day<born.day)
    return dict(id=animal.id,name=animal.name,siniiga=animal.siniiga,sex=animal.sex,birth_date=born,
        age_months=months,status=status,exit_date=exit_date,group=group,
        breeds=[dict(name=s.breed.name,percentage=str(s.percentage)) for s in sorted(animal.breeds,key=lambda s:s.breed.name.casefold())])


@router.get('/animals/options')
def options(actor=Depends(require_permission('animals:read')),db: Session=Depends(get_db,scope='function')):
    return dict(breeds=[dict(id=b.id,name=b.name) for b in db.scalars(select(BreedModel).order_by(BreedModel.name))],
                groups=[dict(id=g.id,name=g.name) for g in db.scalars(select(AnimalGroupModel).order_by(AnimalGroupModel.name))])


@router.get('/animals')
def animals(f=Depends(filters),offset: int=Query(0,ge=0),limit: int=Query(25,ge=1,le=100),
            actor=Depends(require_permission('animals:read')),db: Session=Depends(get_db,scope='function')):
    today=date.today();query=query_for(f,today)
    counts=dict(db.execute(query.with_only_columns(AnimalModel.sex,func.count()).group_by(AnimalModel.sex)).all())
    rows=db.scalars(query.order_by(AnimalModel.name,AnimalModel.siniiga,AnimalModel.id).offset(offset).limit(limit))
    return dict(as_of=today,items=[response(db,a,today) for a in rows],total=sum(counts.values()),female=counts.get('female',0),male=counts.get('male',0),offset=offset,limit=limit)


def text_cell(cell,value):
    cell.value=str(value or '')
    cell.data_type='s'  # Names and identifiers must never become Excel formulas.


@router.get('/animals/export')
def export(f=Depends(filters),actor=Depends(require_permission('animals:read')),db: Session=Depends(get_db,scope='function')):
    today=date.today();workbook=Workbook();sheet=workbook.active;sheet.title='Animales'
    sheet.append(['Pelifolk · Reporte de animales']);sheet.append(['Fecha de consulta',today]);sheet.cell(2,2).number_format='dd/mm/yyyy'
    labels={'active':'Activos','exited':'Bajas','all':'Todos'}
    sheet.append(['Filtros',f"Estado: {labels[f['status']]} · Sexo: { {'female':'Hembra','male':'Macho'}.get(f['sex'],'Todos') }"])
    for label,value in [('Búsqueda',f['q']),('Raza',db.get(BreedModel,f['breed_id']).name if f['breed_id'] and db.get(BreedModel,f['breed_id']) else ''),('Lote actual',db.get(AnimalGroupModel,f['group_id']).name if f['group_id'] and db.get(AnimalGroupModel,f['group_id']) else '')]:
        sheet.append([label]);text_cell(sheet.cell(sheet.max_row,2),value)
    sheet.append([])
    sheet.append(['SINIIGA','Nombre','Sexo','Nacimiento','Edad (meses)','Estado','Fecha de salida','Razas (%)','Lote actual']);header=sheet.max_row
    exits={'sale':'Venta','death':'Muerte','slaughter':'Sacrificio','transfer':'Traspaso','theft':'Robo','other':'Otra salida','active':'Activo'}
    total=0
    for a in db.scalars(query_for(f,today).order_by(AnimalModel.name,AnimalModel.siniiga,AnimalModel.id)):
        row=response(db,a,today);sheet.append([None,None,'Hembra' if row['sex']=='female' else 'Macho',row['birth_date'],row['age_months'],exits[row['status']],row['exit_date'],None,None]);n=sheet.max_row
        for col,value in [(1,row['siniiga']),(2,row['name']),(8,', '.join(s['name']+' '+s['percentage']+'%' for s in row['breeds'])),(9,row['group']['name'] if row['group'] else '')]:text_cell(sheet.cell(n,col),value)
        for col in (4,7):sheet.cell(n,col).number_format='dd/mm/yyyy'
        total+=1
    sheet.auto_filter.ref=f'A{header}:I{sheet.max_row}';sheet.freeze_panes=f'A{header+1}'
    for cell in sheet[header]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='24583A')
    for col,width in zip('ABCDEFGHI',[24,26,14,16,16,18,18,40,26]):sheet.column_dimensions[col].width=width
    sheet.append([]);sheet.append(['Total de animales',total])
    stream=BytesIO();workbook.save(stream)
    return Response(stream.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="pelifolk-animales-{today}.xlsx"','Cache-Control':'no-store'})
